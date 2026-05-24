"""Tests for the upload orchestration service."""

import asyncio
import io
import os
import uuid
from collections.abc import AsyncIterator
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from storage3.utils import StorageException

from app.core.config import get_settings
from app.models.document import DocTypeSource, Document, DocumentStatus
from app.services.ingestion import upload
from app.services.ingestion.antivirus import ScanResult
from app.services.ingestion.archive import ZipBombError
from app.services.ingestion.upload import UploadValidationError, ValidationOutcome

SETTINGS = get_settings()
_BYTES_PER_MB = 1024 * 1024


# --------------------------------------------------------------------------- #
# File-fixture builders (valid, parseable office files)
# --------------------------------------------------------------------------- #
def _pdf_bytes(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = DocxDocument()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_OLE2_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32


# --------------------------------------------------------------------------- #
# Fakes for httpx streaming and async DB sessions
# --------------------------------------------------------------------------- #
class _FakeResponse:
    def __init__(self, data: bytes, *, raise_error: bool = False) -> None:
        self._data = data
        self._raise = raise_error

    def raise_for_status(self) -> None:
        if self._raise:
            raise httpx.HTTPError("download failed")

    async def aiter_bytes(self, chunk_size: int) -> AsyncIterator[bytes]:
        for index in range(0, len(self._data), chunk_size):
            yield self._data[index : index + chunk_size]


class _FakeStreamCtx:
    def __init__(self, response: _FakeResponse) -> None:
        self._response = response

    async def __aenter__(self) -> _FakeResponse:
        return self._response

    async def __aexit__(self, *_: object) -> bool:
        return False


class _FakeHttpClient:
    def __init__(self, response: _FakeResponse) -> None:
        self._response = response

    async def __aenter__(self) -> "_FakeHttpClient":
        return self

    async def __aexit__(self, *_: object) -> bool:
        return False

    def stream(self, _method: str, _url: str) -> _FakeStreamCtx:
        return _FakeStreamCtx(self._response)


class _FakeSessionCtx:
    def __init__(self, session: MagicMock) -> None:
        self._session = session

    async def __aenter__(self) -> MagicMock:
        return self._session

    async def __aexit__(self, *_: object) -> bool:
        return False


def _storage_with_proxy() -> tuple[MagicMock, MagicMock]:
    storage = MagicMock()
    proxy = MagicMock()
    proxy.create_signed_url = AsyncMock(return_value={"signedURL": "https://signed"})
    proxy.create_signed_upload_url = AsyncMock(
        return_value={"signed_url": "https://up", "token": "tok", "path": "p"}
    )
    proxy.remove = AsyncMock(return_value=[])
    proxy.info = AsyncMock()
    storage.from_ = MagicMock(return_value=proxy)
    return storage, proxy


def _make_document(**overrides: object) -> Document:
    defaults: dict[str, object] = {
        "id": uuid.uuid4(),
        "project_id": uuid.uuid4(),
        "original_filename": "01_Spec_Rev01.pdf",
        "storage_bucket": "rfq-documents",
        "storage_path": "proj/doc.pdf",
        "content_type": "application/pdf",
        "size_bytes": 1000,
        "status": DocumentStatus.PROCESSING,
        "doc_type_source": DocTypeSource.AUTO,
    }
    defaults.update(overrides)
    return Document(**defaults)


# --------------------------------------------------------------------------- #
# Pure helpers
# --------------------------------------------------------------------------- #
def test_plan_storage_path() -> None:
    project_id = uuid.uuid4()
    document_id = uuid.uuid4()
    assert upload.plan_storage_path(project_id, document_id, ".pdf") == (
        f"{project_id}/{document_id}.pdf"
    )


def test_validate_upload_request_success() -> None:
    ext = upload.validate_upload_request(
        filename="01_Spec.pdf",
        declared_size=1000,
        existing_count=0,
        existing_total_bytes=0,
        settings=SETTINGS,
    )
    assert ext == ".pdf"


def test_validate_upload_request_bad_extension() -> None:
    with pytest.raises(UploadValidationError) as exc:
        upload.validate_upload_request(
            filename="notes.txt",
            declared_size=10,
            existing_count=0,
            existing_total_bytes=0,
            settings=SETTINGS,
        )
    assert exc.value.status_code == 415


@pytest.mark.parametrize("declared_size", [0, -5, (10**9) * 10])
def test_validate_upload_request_bad_size(declared_size: int) -> None:
    with pytest.raises(UploadValidationError) as exc:
        upload.validate_upload_request(
            filename="01_Spec.pdf",
            declared_size=declared_size,
            existing_count=0,
            existing_total_bytes=0,
            settings=SETTINGS,
        )
    assert exc.value.status_code == 413


def test_validate_upload_request_too_many_files() -> None:
    with pytest.raises(UploadValidationError) as exc:
        upload.validate_upload_request(
            filename="01_Spec.pdf",
            declared_size=10,
            existing_count=SETTINGS.MAX_FILES_PER_PROJECT,
            existing_total_bytes=0,
            settings=SETTINGS,
        )
    assert exc.value.status_code == 409


def test_validate_upload_request_total_size_exceeded() -> None:
    with pytest.raises(UploadValidationError) as exc:
        upload.validate_upload_request(
            filename="01_Spec.pdf",
            declared_size=10,
            existing_count=1,
            existing_total_bytes=SETTINGS.MAX_PROJECT_TOTAL_SIZE_MB * _BYTES_PER_MB,
            settings=SETTINGS,
        )
    assert exc.value.status_code == 413


@pytest.mark.parametrize(
    ("info", "expected"),
    [
        ({"size": 5}, 5),
        ({"contentLength": 7}, 7),
        ({"metadata": {"size": 9}}, 9),
        ({}, None),
        ({"metadata": "not-a-dict"}, None),
    ],
)
def test_extract_size(info: dict[str, object], expected: int | None) -> None:
    assert upload._extract_size(info) == expected


# --------------------------------------------------------------------------- #
# Async helpers with mocks
# --------------------------------------------------------------------------- #
async def test_get_project_document_usage() -> None:
    db = MagicMock()
    result = MagicMock()
    result.one.return_value = (2, 500)
    db.execute = AsyncMock(return_value=result)
    assert await upload.get_project_document_usage(db, uuid.uuid4()) == (2, 500)


async def test_fetch_object_size() -> None:
    storage, proxy = _storage_with_proxy()
    proxy.info = AsyncMock(return_value={"size": 123})
    assert await upload.fetch_object_size(storage, bucket="b", storage_path="p") == 123


async def test_validate_stored_object_pdf_success() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_pdf_bytes(pages=2))
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    assert outcome.ok is True
    assert outcome.page_count == 2
    assert outcome.sha256 is not None


async def test_validate_stored_object_xlsx_success() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_xlsx_bytes())
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".xlsx", settings=SETTINGS
        )
    assert outcome.ok is True
    assert outcome.sheet_names == ["Sheet"]


async def test_validate_stored_object_docx_success() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_docx_bytes())
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".docx", settings=SETTINGS
        )
    assert outcome.ok is True
    assert outcome.page_count is None
    assert outcome.sheet_names is None


async def test_validate_stored_object_xls_accepted_without_parse() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_OLE2_BYTES)
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".xls", settings=SETTINGS
        )
    assert outcome.ok is True
    assert outcome.page_count is None


async def test_validate_stored_object_magic_mismatch() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(b"definitely not a pdf")
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    assert outcome.ok is False
    assert "not a valid" in (outcome.failure_reason or "")


async def test_validate_stored_object_parse_failure() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(b"PK\x03\x04 not really a zip")
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".xlsx", settings=SETTINGS
        )
    assert outcome.ok is False
    assert "parsed" in (outcome.failure_reason or "")


async def test_validate_stored_object_malware_is_permanent() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_pdf_bytes(pages=1))
    with (
        patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)),
        patch.object(
            upload, "scan_file", AsyncMock(return_value=ScanResult(clean=False, signature="Eicar"))
        ),
    ):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    assert outcome.ok is False
    assert outcome.transient is False  # an infected file is a permanent failure
    assert "Malware" in (outcome.failure_reason or "")


async def test_validate_stored_object_scanner_error_is_transient() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_pdf_bytes(pages=1))
    with (
        patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)),
        patch.object(
            upload, "scan_file", AsyncMock(return_value=ScanResult(clean=False, error="refused"))
        ),
    ):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    assert outcome.ok is False
    assert outcome.transient is True  # fail closed: an unreachable scanner is retryable
    assert "Antivirus" in (outcome.failure_reason or "")


async def test_validate_stored_object_zip_bomb_is_permanent() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(_xlsx_bytes())
    with (
        patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)),
        patch.object(
            upload,
            "assert_zip_within_limits",
            MagicMock(side_effect=ZipBombError("archive expands too much")),
        ),
    ):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".xlsx", settings=SETTINGS
        )
    assert outcome.ok is False
    assert outcome.transient is False  # a zip bomb is a permanent content failure
    assert "expands too much" in (outcome.failure_reason or "")


async def test_download_object_to_tempfile_streams_to_disk() -> None:
    storage, _ = _storage_with_proxy()
    payload = _pdf_bytes(pages=1)
    response = _FakeResponse(payload)
    with patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)):
        downloaded = await upload.download_object_to_tempfile(
            storage, bucket="b", storage_path="p", settings=SETTINGS, suffix=".pdf"
        )
    try:
        assert downloaded.size_bytes == len(payload)
        assert downloaded.head.startswith(b"%PDF-")
        assert downloaded.sha256 is not None
        assert os.path.exists(downloaded.path)
    finally:
        os.unlink(downloaded.path)


async def test_download_object_to_tempfile_cleans_up_on_error() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(b"", raise_error=True)
    captured: list[str] = []
    real_mkstemp = upload.tempfile.mkstemp

    def _spy_mkstemp(*args: object, **kwargs: object) -> tuple[int, str]:
        fd, path = real_mkstemp(*args, **kwargs)
        captured.append(path)
        return fd, path

    with (
        patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)),
        patch.object(upload.tempfile, "mkstemp", _spy_mkstemp),
        pytest.raises(httpx.HTTPError),
    ):
        await upload.download_object_to_tempfile(
            storage, bucket="b", storage_path="p", settings=SETTINGS, suffix=".pdf"
        )
    assert captured  # a temp file was created
    assert not os.path.exists(captured[0])  # and cleaned up on failure


async def test_validate_stored_object_download_error_is_transient() -> None:
    storage, _ = _storage_with_proxy()
    response = _FakeResponse(b"", raise_error=True)
    with (
        patch.object(upload.httpx, "AsyncClient", lambda *a, **k: _FakeHttpClient(response)),
        patch.object(upload.anyio, "sleep", AsyncMock()),
    ):
        outcome = await upload.validate_stored_object(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    assert outcome.ok is False
    assert outcome.transient is True  # a download blip must NOT permanently fail the doc
    assert "download" in (outcome.failure_reason or "")


async def test_download_with_retries_succeeds_after_transient_blip() -> None:
    storage, _ = _storage_with_proxy()
    payload = _pdf_bytes(pages=1)
    responses = [_FakeResponse(b"", raise_error=True), _FakeResponse(payload)]

    def _client(*_a: object, **_k: object) -> _FakeHttpClient:
        return _FakeHttpClient(responses.pop(0))

    with (
        patch.object(upload.httpx, "AsyncClient", _client),
        patch.object(upload.anyio, "sleep", AsyncMock()),
    ):
        downloaded = await upload._download_with_retries(
            storage, bucket="b", storage_path="p", ext=".pdf", settings=SETTINGS
        )
    try:
        assert downloaded.size_bytes == len(payload)
    finally:
        os.unlink(downloaded.path)


# --------------------------------------------------------------------------- #
# Background validation + purge
# --------------------------------------------------------------------------- #
async def test_run_document_validation_success() -> None:
    document = _make_document(status=DocumentStatus.PROCESSING)
    session = MagicMock()
    session.get = AsyncMock(return_value=document)
    session.commit = AsyncMock()
    storage, _ = _storage_with_proxy()
    outcome = ValidationOutcome(
        ok=True, failure_reason=None, page_count=3, sheet_names=None, sha256="abc", size_bytes=55
    )
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "validate_stored_object", AsyncMock(return_value=outcome)),
    ):
        await upload.run_document_validation(storage, document.id)
    assert document.status == DocumentStatus.READY
    assert document.page_count == 3
    assert document.sha256 == "abc"
    assert document.size_bytes == 55
    session.commit.assert_awaited_once()


async def test_run_document_validation_failure_removes_object() -> None:
    document = _make_document(status=DocumentStatus.PROCESSING)
    session = MagicMock()
    session.get = AsyncMock(return_value=document)
    session.commit = AsyncMock()
    storage, proxy = _storage_with_proxy()
    outcome = ValidationOutcome(
        ok=False,
        failure_reason="bad file",
        page_count=None,
        sheet_names=None,
        sha256=None,
        size_bytes=None,
    )
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "validate_stored_object", AsyncMock(return_value=outcome)),
    ):
        await upload.run_document_validation(storage, document.id)
    assert document.status == DocumentStatus.FAILED
    assert document.failure_reason == "bad file"
    proxy.remove.assert_awaited_once()


async def test_run_document_validation_skips_missing_document() -> None:
    session = MagicMock()
    session.get = AsyncMock(return_value=None)
    session.commit = AsyncMock()
    storage, _ = _storage_with_proxy()
    with patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)):
        await upload.run_document_validation(storage, uuid.uuid4())
    session.commit.assert_not_awaited()


async def test_run_document_validation_storage_exception_is_transient() -> None:
    """A storage-layer error is retryable: leave the doc processing, keep its bytes."""
    document = _make_document(status=DocumentStatus.PROCESSING)
    session = MagicMock()
    session.get = AsyncMock(return_value=document)
    session.commit = AsyncMock()
    storage, proxy = _storage_with_proxy()
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(
            upload, "validate_stored_object", AsyncMock(side_effect=StorageException("gone"))
        ),
    ):
        await upload.run_document_validation(storage, document.id)
    assert document.status == DocumentStatus.PROCESSING
    proxy.remove.assert_not_awaited()
    session.commit.assert_not_awaited()


async def test_run_document_validation_transient_leaves_processing() -> None:
    """A transient validation outcome must not fail the doc or delete the object."""
    document = _make_document(status=DocumentStatus.PROCESSING)
    session = MagicMock()
    session.get = AsyncMock(return_value=document)
    session.commit = AsyncMock()
    storage, proxy = _storage_with_proxy()
    outcome = ValidationOutcome(
        ok=False,
        transient=True,
        failure_reason="blip",
        page_count=None,
        sheet_names=None,
        sha256=None,
        size_bytes=None,
    )
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "validate_stored_object", AsyncMock(return_value=outcome)),
    ):
        await upload.run_document_validation(storage, document.id)
    assert document.status == DocumentStatus.PROCESSING
    proxy.remove.assert_not_awaited()
    session.commit.assert_not_awaited()


def _purge_session(documents: list[Document]) -> MagicMock:
    session = MagicMock()
    result = MagicMock()
    scalars = MagicMock()
    scalars.all.return_value = documents
    result.scalars.return_value = scalars
    session.execute = AsyncMock(return_value=result)
    session.delete = AsyncMock()
    session.commit = AsyncMock()
    return session


async def test_purge_stale_pending_documents_deletes_and_removes() -> None:
    document = _make_document(status=DocumentStatus.PENDING)
    session = _purge_session([document])
    storage, proxy = _storage_with_proxy()
    with patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)):
        await upload.purge_stale_pending_documents(
            storage, project_id=uuid.uuid4(), settings=SETTINGS
        )
    session.delete.assert_awaited_once_with(document)
    proxy.remove.assert_awaited_once()
    session.commit.assert_awaited_once()


async def test_purge_swallows_storage_errors() -> None:
    document = _make_document(status=DocumentStatus.PENDING)
    session = _purge_session([document])
    storage, proxy = _storage_with_proxy()
    proxy.remove = AsyncMock(side_effect=StorageException("nope"))
    with patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)):
        await upload.purge_stale_pending_documents(
            storage, project_id=uuid.uuid4(), settings=SETTINGS
        )
    session.delete.assert_awaited_once_with(document)
    session.commit.assert_awaited_once()


async def test_purge_never_raises_on_db_error() -> None:
    session = MagicMock()
    session.execute = AsyncMock(side_effect=RuntimeError("db down"))
    storage, _ = _storage_with_proxy()
    with patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)):
        # Must not raise.
        await upload.purge_stale_pending_documents(
            storage, project_id=uuid.uuid4(), settings=SETTINGS
        )


# --------------------------------------------------------------------------- #
# Stuck-processing recovery
# --------------------------------------------------------------------------- #
def _ids_session(doc_ids: list[uuid.UUID]) -> MagicMock:
    session = MagicMock()
    result = MagicMock()
    result.scalars.return_value.all.return_value = doc_ids
    session.execute = AsyncMock(return_value=result)
    return session


async def test_recover_stuck_processing_redrives_each_document() -> None:
    doc_ids = [uuid.uuid4(), uuid.uuid4()]
    session = _ids_session(doc_ids)
    storage, _ = _storage_with_proxy()
    redrive = AsyncMock()
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "run_document_validation", redrive),
    ):
        await upload.recover_stuck_processing_documents(storage, settings=SETTINGS)
    assert redrive.await_count == 2


async def test_recover_stuck_processing_noop_when_none() -> None:
    session = _ids_session([])
    storage, _ = _storage_with_proxy()
    redrive = AsyncMock()
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "run_document_validation", redrive),
    ):
        await upload.recover_stuck_processing_documents(storage, settings=SETTINGS)
    redrive.assert_not_awaited()


async def test_recover_stuck_processing_swallows_query_error() -> None:
    session = MagicMock()
    session.execute = AsyncMock(side_effect=RuntimeError("db down"))
    storage, _ = _storage_with_proxy()
    with patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)):
        # Must not raise — recovery can never block startup.
        await upload.recover_stuck_processing_documents(storage, settings=SETTINGS)


async def test_recover_stuck_processing_continues_past_failures() -> None:
    doc_ids = [uuid.uuid4(), uuid.uuid4()]
    session = _ids_session(doc_ids)
    storage, _ = _storage_with_proxy()
    redrive = AsyncMock(side_effect=[RuntimeError("boom"), None])
    with (
        patch.object(upload, "AsyncSessionLocal", lambda: _FakeSessionCtx(session)),
        patch.object(upload, "run_document_validation", redrive),
    ):
        await upload.recover_stuck_processing_documents(storage, settings=SETTINGS)
    assert redrive.await_count == 2  # the first failing does not stop the second


# --------------------------------------------------------------------------- #
# Periodic recovery loop
# --------------------------------------------------------------------------- #
async def test_run_recovery_loop_sweeps_then_sleeps() -> None:
    storage, _ = _storage_with_proxy()
    sweep = AsyncMock()
    sleep = AsyncMock(side_effect=asyncio.CancelledError)  # stop after the first iteration
    with (
        patch.object(upload, "recover_stuck_processing_documents", sweep),
        patch.object(upload.anyio, "sleep", sleep),
        pytest.raises(asyncio.CancelledError),
    ):
        await upload.run_recovery_loop(storage, settings=SETTINGS)
    sweep.assert_awaited_once()
    sleep.assert_awaited_once()


async def test_run_recovery_loop_survives_sweep_error() -> None:
    storage, _ = _storage_with_proxy()
    sweep = AsyncMock(side_effect=RuntimeError("boom"))
    sleep = AsyncMock(side_effect=asyncio.CancelledError)
    with (
        patch.object(upload, "recover_stuck_processing_documents", sweep),
        patch.object(upload.anyio, "sleep", sleep),
        pytest.raises(asyncio.CancelledError),
    ):
        await upload.run_recovery_loop(storage, settings=SETTINGS)
    sweep.assert_awaited_once()  # error swallowed; the loop still reached the sleep
    sleep.assert_awaited_once()
